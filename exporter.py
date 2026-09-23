from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    ListFlowable,
    ListItem
)
from reportlab.lib.styles import getSampleStyleSheet

from docx import Document


class MeetingExporter:

    @staticmethod
    def export_pdf(state, path):

        document = SimpleDocTemplate(
            path,
            pagesize=A4,
            rightMargin=50,
            leftMargin=50,
            topMargin=50,
            bottomMargin=50
        )

        styles = getSampleStyleSheet()

        title_style = styles["Title"]
        heading_style = styles["Heading2"]
        body_style = styles["BodyText"]

        story = []

        meeting_name = (
            state.meeting_name
            or "Meeting Report"
        )

        story.append(
            Paragraph(
                meeting_name,
                title_style
            )
        )

        story.append(Spacer(1, 20))

        if state.participants:

            story.append(
                Paragraph(
                    "Participants",
                    heading_style
                )
            )

            for participant_id, name in state.participants.items():

                text = (
                    f"{name} "
                    f"({participant_id})"
                )

                story.append(
                    Paragraph(
                        text,
                        body_style
                    )
                )

            story.append(Spacer(1, 15))

        if state.summary:

            story.append(
                Paragraph(
                    "Summary",
                    heading_style
                )
            )

            story.append(
                Paragraph(
                    state.summary,
                    body_style
                )
            )

            story.append(Spacer(1, 15))

        MeetingExporter._add_list_section(
            story,
            "Topics",
            state.topics,
            heading_style,
            body_style
        )

        MeetingExporter._add_list_section(
            story,
            "Decisions",
            state.decisions,
            heading_style,
            body_style
        )

        MeetingExporter._add_list_section(
            story,
            "Action Items",
            state.action_items,
            heading_style,
            body_style
        )

        MeetingExporter._add_list_section(
            story,
            "Open Questions",
            state.open_questions,
            heading_style,
            body_style
        )

        document.build(story)

    @staticmethod
    def _add_list_section(
        story,
        title,
        items,
        heading_style,
        body_style
    ):

        if not items:
            return

        story.append(
            Paragraph(
                title,
                heading_style
            )
        )

        flowable_items = []

        for item in items:

            flowable_items.append(
                ListItem(
                    Paragraph(
                        item,
                        body_style
                    )
                )
            )

        story.append(
            ListFlowable(
                flowable_items,
                bulletType="bullet"
            )
        )

        story.append(Spacer(1, 15))

    @staticmethod
    def export_docx(state, path):

        document = Document()

        meeting_name = (
            state.meeting_name
            or "Meeting Report"
        )

        document.add_heading(
            meeting_name,
            level=0
        )

        if state.participants:

            document.add_heading(
                "Participants",
                level=1
            )

            for participant_id, name in state.participants.items():

                document.add_paragraph(
                    f"{name} ({participant_id})"
                )

        if state.summary:

            document.add_heading(
                "Summary",
                level=1
            )

            document.add_paragraph(
                state.summary
            )

        MeetingExporter._add_docx_list(
            document,
            "Topics",
            state.topics
        )

        MeetingExporter._add_docx_list(
            document,
            "Decisions",
            state.decisions
        )

        MeetingExporter._add_docx_list(
            document,
            "Action Items",
            state.action_items
        )

        MeetingExporter._add_docx_list(
            document,
            "Open Questions",
            state.open_questions
        )

        document.save(path)

    @staticmethod
    def _add_docx_list(
        document,
        title,
        items
    ):

        if not items:
            return

        document.add_heading(
            title,
            level=1
        )

        for item in items:

            document.add_paragraph(
                item,
                style="List Bullet"
            )